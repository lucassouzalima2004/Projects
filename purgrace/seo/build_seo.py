"""Builds the SEO spreadsheet and the Shopify import files for PurGrace.

Reads the store's public catalogue (dev/fixtures/purgrace-store.json) and the SEO texts the live store
shows today (live-seo.json, from dev/fetch-live-seo.mjs), and writes:
  PurGrace-SEO.xlsx                          review sheet (products, collections, issues, rings)
  shopify-import-0-teste-2-produtos.csv      import 1 for two products only, to test first
  shopify-import-1-titulos-tipos-tags.csv    titles, types, tags, category, SKU and MPN (no SEO fields)
  shopify-import-2-descricoes.csv            new descriptions for every product, from descricoes.json
  shopify-import-3-seo-onde-falta.csv        SEO title and description only where the product has none today
  shopify-import-desfazer.csv                puts titles, types, tags, SKU, MPN and descriptions back as they were

Usage: python3 build_seo.py   (needs openpyxl: pip install openpyxl)
"""
import csv
import html
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
FIXTURE = HERE.parent / "dev" / "fixtures" / "purgrace-store.json"
LIVE = HERE / "live-seo.json"
# New description per product (lead, size, stones, details), written from the title and the facts in today's text
DESCRIPTIONS = HERE / "descricoes.json"
STORE = "https://purgrace.com.au"

# Product types (used by filters, Google and automated collections) and their Shopify categories
TYPES = {
    "E": ("Earrings", "Apparel & Accessories > Jewelry > Earrings", "earrings"),
    "N": ("Necklaces", "Apparel & Accessories > Jewelry > Necklaces", "necklace"),
    "P": ("Pendants", "Apparel & Accessories > Jewelry > Charms & Pendants > Pendants", "pendant"),
    "R": ("Rings", "Apparel & Accessories > Jewelry > Rings", "ring"),
    "B": ("Bracelets", "Apparel & Accessories > Jewelry > Bracelets", "bracelet"),
    "A": ("Anklets", "Apparel & Accessories > Jewelry > Anklets", None),
    "S": ("Jewellery Sets", "Apparel & Accessories > Jewelry > Jewelry Sets", "set"),
}

# Clean title, type and theme tags for every product, in catalogue order (checked against the handle below).
# Titles: material first, no product code, sizes only where they tell two similar pieces apart.
CURATED = [
    ("18k-gold-plated-saint-benedict-medal", "18K Gold-Plated Saint Benedict Medal Pendant", "P", "faith, saint benedict, medal"),
    ("pendant-in-the-shape-of-a-girl", "18K Gold-Plated Pavé Girl Pendant", "P", "family, girl"),
    ("heart-choker", "18K Gold-Plated Heart Link Necklace, 42 cm", "N", "love, heart"),
    ("girl-pendant-plated", "18K Gold-Plated Girl Pendant with Rhodium Plating", "P", "family, girl, two-tone"),
    ("butterfly-earrings-with-zirconia-527074", "18K Gold-Plated Butterfly Earrings with Navette Cubic Zirconia", "E", "butterfly"),
    ("necklace-53069500", "18K Gold-Plated Fine Chain Necklace, 50 cm", "N", "chain"),
    ("butterfly-set-0003", "18K Gold-Plated Open Butterfly Necklace and Earrings Set", "S", "butterfly, gift set"),
    ("butterfly-pendant-choker", "18K Gold-Plated Open Butterfly Pendant Necklace", "N", "butterfly"),
    ("textured-cross-pendant", "18K Gold-Plated Textured Cross Pendant", "P", "faith, cross, unisex"),
    ("mens-18k-gold-plated-chain", "Men's 18K Gold-Plated Chain Necklace, 60 cm", "N", "men, chain"),
    ("dragonfly-anklet", "18K Gold-Plated Dragonfly Anklet, 25 cm", "A", "nature, dragonfly"),
    ("choker-with-baroque-pearl", "18K Gold-Plated Baroque Pearl Pendant Necklace", "N", "pearl"),
    ("holy-spirit-earring", "18K Gold-Plated Holy Spirit Dove Earrings", "E", "faith, holy spirit"),
    ("hoop-earring-5236460000", "18K Gold-Plated Triple Hoop Earrings", "E", "hoops"),
    ("four-leaf-clover-earring", "18K Gold-Plated Four-Leaf Clover Earrings", "E", "luck, clover"),
    ("clip-on-ear-cuff", "18K Gold-Plated Clip-On Ear Cuff", "E", "ear cuff"),
    ("hoop-earrings-with-cubic-zirconia-5272370006", "18K Gold-Plated Hoop Earrings with Pear Cubic Zirconia Drops", "E", "hoops"),
    ("cross-hoop-earring", "18K Gold-Plated Hoop Earrings with Cross Drops", "E", "faith, cross, hoops"),
    ("flower-earring-set", "18K Gold-Plated Flower Stud Earring Set", "E", "nature, flower, studs"),
    ("crosses-scapular", "Men's 18K Gold-Plated Cross Scapular Necklace, 68 cm", "N", "faith, cross, men, scapular"),
    ("unisex-scapular", "18K Gold-Plated Our Lady of Mount Carmel Scapular Necklace, 49 cm", "N", "faith, our lady, scapular, unisex"),
    ("necklace-with-lock-chain-5316384200", "18K Gold-Plated Fine Lock Chain Necklace, 42 cm", "N", "chain"),
    ("venetian-choker-plated", "18K Gold-Plated Venetian Chain Necklace, 42 cm", "N", "chain"),
    ("butterfly-earring-with-openwork-design", "18K Gold-Plated Openwork Butterfly Earrings", "E", "butterfly, studs"),
    ("boy-pendant-with-rhodium", "18K Gold-Plated Boy Pendant with Rhodium Plating", "P", "family, boy, two-tone"),
    ("half-hoop-earring-set", "18K Gold-Plated Half Hoop Earring Set, 2 Pairs", "E", "hoops, ana hickmann"),
    ("butterfly-earring-plated-in-18k-gold-52709200", "18K Gold-Plated Openwork Butterfly Studs, 1.2 cm", "E", "butterfly, studs"),
    ("our-lady-of-aparecida", "18K Gold-Plated Our Lady of Aparecida Pendant with Blue Cubic Zirconia", "P", "faith, our lady, aparecida"),
    ("set-gold-plated-butterfly", "18K Gold-Plated Mini Butterfly Pendant and Earrings Set", "S", "butterfly, gift set"),
    ("solitaire-earring-with-crystals-5201850006", "18K Gold-Plated Bezel Crystal Stud Earrings, 4 mm", "E", "studs, crystal"),
    ("cable-chain-necklace-5304325000", "18K Gold-Plated Cable Chain Necklace, 50 cm", "N", "chain"),
    ("pet-mom", "18K Gold-Plated Pet Mum Double-Sided Medal Pendant", "P", "family, pets, medal"),
    ("holy-spirit-pendant-5426970000", "18K Gold-Plated Holy Spirit Heart Pendant", "P", "faith, holy spirit, heart"),
    ("faith-pendant", "18K Gold-Plated Fé (Faith) Pendant with Cubic Zirconia", "P", "faith"),
    ("blessed-sacrament", "18K Gold-Plated Blessed Sacrament Pendant with Resin Stone", "P", "faith"),
    ("holy-spirit-pendant-5419330006", "18K Gold-Plated Round Holy Spirit Pendant with Cubic Zirconia", "P", "faith, holy spirit"),
    ("cable-chain-necklace-5309604500", "18K Gold-Plated Cable Chain Necklace, 45 cm", "N", "chain"),
    ("gargantilha", "18K Gold-Plated Curb Chain Necklace, 42 cm", "N", "chain"),
    ("solitaire-earring-with-crystals-5231910006", "18K Gold-Plated Crystal Solitaire Stud Earrings", "E", "studs, crystal"),
    ("solitaire-heart-earrings", "Children's 18K Gold-Plated Heart Stud Earrings with Cubic Zirconia", "E", "kids, heart, studs"),
    ("childrens-earrings", "Children's 18K Gold-Plated Ball Drop Earrings", "E", "kids"),
    ("set-open-heart", "18K Gold-Plated Open Heart Necklace and Earrings Set", "S", "love, heart, gift set"),
    ("hoop-earrings-5240140000", "18K Gold-Plated Textured Hoop Earrings, 1.3 cm", "E", "hoops"),
    ("bow-earrings", "Children's 18K Gold-Plated Bow Stud Earrings", "E", "kids, bow, studs"),
    ("hoop-earring-5224420000", "18K Gold-Plated Wide Huggie Hoop Earrings, 1.3 cm", "E", "hoops"),
    ("earring-with-crystals", "18K Gold-Plated Crystal Ear Climber Earrings", "E", "ear climber, crystal"),
    ("oval-hoop-earrings", "18K Gold-Plated Oval Hoop Earrings, 1.8 cm", "E", "hoops"),
    ("necklace-53030700", "18K Gold-Plated Hammered Chain Necklace, 42 cm", "N", "chain"),
    ("pingente-solitario", "18K Gold-Plated Solitaire Pendant Necklace, 40 cm", "N", "giovanna antonelli"),
    ("heart-pendant-5429910000", "18K Gold-Plated Puffed Heart Pendant", "P", "love, heart"),
    ("52080700", "18K Gold-Plated Classic Hoop Earrings, 1.8 cm", "E", "hoops"),
    ("hoop-earrings-5209010000", "18K Gold-Plated Smooth Hoop Earrings, 1.6 cm", "E", "hoops"),
    ("rhodium-plating-52016400", "18K Gold-Plated Butterfly Studs with Rhodium Plating", "E", "butterfly, studs, two-tone"),
    ("pearls-and-crystals", "18K Gold-Plated Pearl and Crystal Stud Earrings", "E", "pearl, crystal, studs"),
    ("butterfly-earrings-with-zirconia-52639506", "18K Gold-Plated Mini Butterfly Studs with Cubic Zirconia", "E", "butterfly, studs"),
    ("ear-cuff", "18K Gold-Plated Cubic Zirconia Ear Climber Earrings", "E", "ear climber"),
    ("circulos-entrelacados", "18K Gold-Plated Interlocking Circles Earrings", "E", "studs"),
    ("synthetic-pearls", "18K Gold-Plated Pearl Drop Earrings", "E", "pearl"),
    ("girl-pendant-5406840000", "18K Gold-Plated Girl Face Pendant", "P", "family, girl"),
    ("boy-pendant-5406850000", "18K Gold-Plated Boy Face Pendant", "P", "family, boy"),
    ("choker-with-lock-chain", "18K Gold-Plated Lock Chain Necklace, 42 cm", "N", "chain"),
    ("domed-frieze", "18K Gold-Plated Grooved Heart Pendant", "P", "love, heart"),
    ("butterfly-pendant-5423050000", "18K Gold-Plated Openwork Butterfly Pendant", "P", "butterfly"),
    ("cross-pendant-with-zirconias", "18K Gold-Plated Cross Pendant with Cubic Zirconia", "P", "faith, cross"),
    ("childrens-pendant-boy", "18K Gold-Plated Boy Playing Soccer Pendant", "P", "family, boy"),
    ("butterfly-earring-plated-in-18k-gold-52756400", "18K Gold-Plated Large Openwork Butterfly Earrings", "E", "butterfly"),
    ("two-boys", "18K Gold-Plated Two Boys and Heart Pendant", "P", "family, boy, heart"),
    ("two-girls", "18K Gold-Plated Two Girls and Heart Pendant", "P", "family, girl, heart"),
    ("boy-and-girl-pendant", "18K Gold-Plated Boy and Girl with Heart Pendant", "P", "family, boy, girl, heart"),
    ("heart-shaped-hoop", "18K Gold-Plated Hoop Earrings with Heart Drops", "E", "love, heart, hoops"),
    ("maxi-ring-with-intertwined", "18K Gold-Plated Maxi Wire Ring with Sphere", "R", "statement"),
    ("maxi-triple-band", "18K Gold-Plated Maxi Triple Band Ring with Cubic Zirconia", "R", "statement"),
    ("butterfly-earring-plated-in-18k-gold-52656000", "18K Gold-Plated Mini Openwork Butterfly Studs, 0.9 cm", "E", "butterfly, studs"),
    ("hoop-earrings-with-cubic-zirconia-18k-gold-plated-52677006", "18K Gold-Plated Hoop Earrings with Cubic Zirconia Drops", "E", "hoops"),
    ("heart-earrings-52754600", "18K Gold-Plated Beaded Heart Stud Earrings", "E", "love, heart, studs"),
    ("lotso", "18K Gold-Plated Lotso and Strawberry Bracelet", "B", "kids, disney"),
    ("boy-pendant-with-zirconia", "18K Gold-Plated Boy on a Swing Pendant with Cubic Zirconia", "P", "family, boy"),
    ("starfish", "18K Gold-Plated Starfish Stud Earrings", "E", "nature, starfish, studs"),
    ("skinny-infinity-ring", "18K Gold-Plated Skinny Infinity Ring", "R", "stacking"),
    ("double-band-ring", "18K Gold-Plated Double Band Ring with Spheres and Cubic Zirconia", "R", "stacking"),
    ("skinny-butterfly-ring", "18K Gold-Plated Skinny Butterfly Ring with Cubic Zirconia", "R", "butterfly, stacking"),
    ("customizable", "18K Gold-Plated Personalised Bar Necklace", "N", "personalised"),
    ("brinco-coracao-banhado", "18K Gold-Plated Heart Crystal Earrings", "E", "love, heart, crystal"),
    ("earrings-with-crystals-and-synthetic-pearls", "18K Gold-Plated Crystal and Pearl Drop Earrings", "E", "pearl, crystal"),
    ("curved-earring", "18K Gold-Plated Curved Cubic Zirconia Earrings", "E", "statement"),
    ("pingente-esfera", "18K Gold-Plated Ball Pendant Necklace, 45 cm", "N", ""),
    ("butterfly-earring-with-zirconias", "18K Gold-Plated Pavé Butterfly Stud Earrings", "E", "butterfly, studs"),
    ("triangle-earring-set", "18K Gold-Plated Triangle Stud and Hoop Earring Set", "E", "hoops, studs"),
    ("half-hoop-earrings-527079", "18K Gold-Plated Ribbed Half Hoop Earrings, 1.3 cm", "E", "hoops"),
    ("childrens-solitaire-earring", "Children's 18K Gold-Plated Star Stud Earrings with Crystal", "E", "kids, star, studs"),
    ("wavy-hoop", "18K Gold-Plated Wavy Hoop Earrings with Rhodium Plating", "E", "hoops, two-tone"),
    ("brinco-argola-banhado-a-ouro-18k-com-zirconias-52321406", "18K Gold-Plated Huggie Hoop Earrings with Cubic Zirconia", "E", "hoops"),
    ("hook-earring-with-spheres", "18K Gold-Plated Filigree Ball Drop Earrings", "E", "statement"),
    ("salamander", "18K Gold-Plated Salamander Chain Choker, 38 cm", "N", "chain, choker"),
    ("double-link-bracelet", "925 Sterling Silver Double Link Bracelet, 19 cm", "B", ""),
    ("heart-solitaire-pendant", "925 Sterling Silver Red Heart Pendant with Cubic Zirconia", "P", "love, heart"),
    ("gourmet-link-necklace-830054", "925 Sterling Silver Gourmet Chain Necklace, 50 cm", "N", "chain"),
    ("silver-butterfly-earrings", "925 Sterling Silver Beaded Butterfly Stud Earrings", "E", "butterfly, studs"),
    ("solitaire-teardrop-earring", "925 Sterling Silver Teardrop Cubic Zirconia Stud Earrings", "E", "studs"),
    ("butterfly-earring-in-925", "925 Sterling Silver Butterfly Stud Earrings with Cubic Zirconia", "E", "butterfly, studs"),
    ("ball-earrings", "925 Sterling Silver Ball Stud Earrings, 3 mm", "E", "studs, kids"),
    ("925-sterling-silver-hoop-earrings", "925 Sterling Silver Huggie Hoop Earrings with Cubic Zirconia", "E", "hoops"),
    ("crown-earrings", "Children's 18K Gold-Plated Crown Stud Earrings", "E", "kids, crown, studs"),
    ("bird-earrings", "Children's 18K Gold-Plated Bird Drop Earrings", "E", "kids, nature, bird"),
    ("earrings-with-zirconias-52405106", "18K Gold-Plated Pavé Heart Stud Earrings", "E", "love, heart, studs"),
    ("half-hoop-earring-with-zirconias", "18K Gold-Plated Pavé Half Hoop Earrings", "E", "hoops"),
    ("half-hoop-earrings-52672600", "18K Gold-Plated Smooth Half Hoop Earrings, 1.3 cm", "E", "hoops"),
    ("childrens-bracelet", "Children's 18K Gold-Plated Bead Bracelet, 14 cm", "B", "kids"),
    ("diamond-shaped-chain", "18K Gold-Plated Diamond-Shaped Link Necklace, 42 cm", "N", "chain"),
    ("choker-with-lock-chain-53163800", "18K Gold-Plated Fine Lock Chain Necklace, 42 cm", "N", "chain"),
    ("synthetic-stones-52756627", "18K Gold-Plated Black Butterfly Stud Earrings", "E", "butterfly, studs"),
    ("synthetic-stones-52756506", "18K Gold-Plated White Butterfly Stud Earrings", "E", "butterfly, studs"),
    ("dometer-earring", "18K Gold-Plated Teardrop Stud Earrings", "E", "studs"),
    ("earring-with-sphere", "18K Gold-Plated Ball Stud Earrings, 8 mm", "E", "studs"),
    ("sphere-earrings", "18K Gold-Plated Ball Stud Earrings, 4 mm", "E", "studs"),
    ("openwork-hoop-earrings", "18K Gold-Plated Openwork Hoop Earrings, 1.4 cm", "E", "hoops"),
    ("hoop-earrings-52290600", "18K Gold-Plated Plain Hoop Earrings, 1.7 cm", "E", "hoops"),
    ("half-eternity-ring-in-925", "925 Sterling Silver Half Eternity Ring with Cubic Zirconia", "R", "stacking"),
    ("gourmet-link-choker-830053", "925 Sterling Silver Gourmet Chain Necklace, 42 cm", "N", "chain"),
    ("teardrop-pendant", "925 Sterling Silver Green Teardrop Pendant with Cubic Zirconia", "P", ""),
    ("teardrop-earrings", "925 Sterling Silver Green Teardrop Stud Earrings", "E", "studs"),
    ("solitaire-ring-with-cubic-zirconia", "925 Sterling Silver Solitaire Ring with Cubic Zirconia", "R", ""),
    ("solitaire-ring-with-zirconia", "18K Gold-Plated Solitaire Ring with Cubic Zirconia", "R", ""),
    ("half-eternity-ring-plated", "18K Gold-Plated Half Eternity Ring with Cubic Zirconia", "R", "stacking"),
    ("ring-with-zirconias-511384", "18K Gold-Plated Pavé Disc Ring", "R", ""),
    ("ring-with-crystal-and-zirconias", "18K Gold-Plated Red Crystal Halo Ring", "R", "statement"),
    ("heart-bracelet", "18K Gold-Plated Heart Link Bracelet, 19 cm", "B", "love, heart"),
    ("adjustable", "18K Gold-Plated Y Necklace with Cubic Zirconia", "N", ""),
    ("hoop-earrings-52403200", "18K Gold-Plated Textured Hoop Earrings, 1.1 cm", "E", "hoops"),
    ("hoop-earrings-with-rhodium-plating", "18K Gold-Plated Hoop Earrings with Rhodium Plating, 1.9 cm", "E", "hoops, two-tone"),
    ("round-earrings", "18K Gold-Plated Round Pavé Stud Earrings", "E", "studs"),
    ("solitaire-earring-set", "18K Gold-Plated Solitaire Stud Earring Set, 2 Pairs", "E", "studs"),
    ("arrow-earring", "18K Gold-Plated Arrow and Heart Stud Earring Set", "E", "love, heart, studs"),
    ("ring-51194500", "18K Gold-Plated Wide Twisted Band Ring", "R", "statement"),
    ("teardrop-hoop", "925 Sterling Silver Teardrop Hoop Earrings", "E", "hoops"),
    ("necklace-with-teardrop-pendant", "925 Sterling Silver Teardrop Pendant Necklace, 50 cm", "N", ""),
    ("bracelet-with-crystals", "18K Gold-Plated Crystal Tennis Bracelet", "B", "crystal"),
    ("mother-of-a-boy", "18K Gold-Plated Mother of a Boy Bracelet", "B", "family, mum, boy"),
    ("best-friends", "18K Gold-Plated Best Friends Rainbow Pendants", "P", "friendship"),
    ("square-solitaire", "18K Gold-Plated Blue Square Crystal Stud Earrings", "E", "studs, crystal"),
    ("clip-on-earring", "18K Gold-Plated Green Crystal Clip-On Ear Cuff", "E", "ear cuff, crystal"),
    ("half-hoop-earrings-with-zirconias", "18K Gold-Plated Green Cubic Zirconia Half Hoop Earrings", "E", "hoops"),
    ("ring-with-crystal-513403", "18K Gold-Plated Blue Crystal Double Band Ring", "R", "statement"),
    ("heart-solitaire-ring", "18K Gold-Plated Black Heart Solitaire Ring", "R", "love, heart"),
    ("ring-with-zirconias-and-crystal", "18K Gold-Plated Blue Crystal Ring with Cubic Zirconia", "R", "statement"),
    ("butterfly-bracelet", "18K Gold-Plated Butterfly Bracelet with White Stones", "B", "butterfly"),
    ("bracelet-55127100", "18K Gold-Plated Twisted Oval Link Bracelet", "B", ""),
    ("butterfly-pendant-with-zirconia", "18K Gold-Plated Butterfly Pendant with Purple Cubic Zirconia", "P", "butterfly"),
    ("cinderella-dress", "18K Gold-Plated Cinderella Dress Pendant", "P", "disney, kids"),
    ("venetian-chain", "18K Gold-Plated Venetian Bead Chain Necklace, 80 cm", "N", "chain"),
    ("necklace-with-heart", "18K Gold-Plated Pavé Heart Pendant Necklace, 50 cm", "N", "love, heart"),
    ("piastrine", "18K Gold-Plated Piastrine Chain Necklace, 42 cm", "N", "chain"),
    ("cinderella-shoe", "18K Gold-Plated Cinderella Slipper Earrings", "E", "disney, kids"),
    ("oval-links", "18K Gold-Plated Oval Link Bracelet", "B", ""),
    ("bracelet-with-hearts", "18K Gold-Plated Open Heart Bracelet, 19 cm", "B", "love, heart"),
    ("bracelet-with-balls", "18K Gold-Plated Ball Bracelet", "B", ""),
    ("shape-of-a-boy", "18K Gold-Plated Pavé Boy Pendant", "P", "family, boy"),
    ("groumet", "18K Gold-Plated Gourmet Chain Choker, 35 cm", "N", "chain, choker"),
    ("apple-shape", "Children's 18K Gold-Plated Apple Stud Earrings with Crystal", "E", "kids, studs"),
    ("rhodium-application", "18K Gold-Plated Hoop Earrings with Rhodium Plating, 1.6 cm", "E", "hoops, two-tone"),
    ("openheart", "18K Gold-Plated Open Heart Hoop Earrings", "E", "love, heart, hoops"),
    ("butterfly-earrings-52698900", "18K Gold-Plated Textured Butterfly Stud Earrings", "E", "butterfly, studs"),
]

# Pieces that need a decision before or after the import (product index -> note in Portuguese)
FLAGS = {
    21: "Parece ser o mesmo colar do #109 (mesmo código-base 531638, mesma descrição). Mantenha um só e crie um redirecionamento do outro.",
    109: "Parece ser o mesmo colar do #21 (mesmo código-base 531638, mesma descrição). Mantenha um só e crie um redirecionamento do outro.",
    121: "Anel de prata 925 com o código 510139, que é o mesmo do anel banhado a ouro (#122). Códigos de prata da Rommanel começam com 8: confira o código certo.",
    143: "A descrição diz 'Size: 16 (BR)', mas o anúncio oferece 7 tamanhos (N a Z+1). Confira quais tamanhos você tem de verdade.",
    144: "A descrição diz 'Size: 14(BR)', mas o anúncio oferece 7 tamanhos (N a Z+1). Confira quais tamanhos você tem de verdade.",
    70: "Anel sem opção de tamanho: o código termina em ...1400, ou seja aro 14 (confira na peça). Informe o tamanho no anúncio.",
    78: "Anel sem opção de tamanho: o código termina em ...1400, ou seja aro 14 (confira na peça). Informe o tamanho no anúncio.",
    79: "Anel sem opção de tamanho: o código termina em ...1206, ou seja aro 12 (confira na peça). Informe o tamanho no anúncio.",
    80: "Anel sem opção de tamanho: o código termina em ...1606, ou seja aro 16 (confira na peça). Informe o tamanho no anúncio.",
    71: "Anel sem opção de tamanho e com código curto (513639), sem o aro: veja o aro na etiqueta da peça e informe no anúncio.",
    93: "A descrição diz 'Medical grade' sem comprovação: a descrição nova da importação 2 não traz essa frase.",
    54: "O título e a descrição para o Google de hoje falam de anel ('Butterfly Ring for Girls'), mas a peça é um par de brincos de borboleta. Corrija em Search engine listing: as importações não mexem nesses textos.",
    150: "A descrição diz 'Triple gold finish (24k, 18k & 22k)'. A Rommanel não publica essa divisão: a descrição nova da importação 2 diz 'three layers of gold'.",
    151: "A descrição diz 'Triple gold finish (24k, 18k & 22k)'. A Rommanel não publica essa divisão: a descrição nova da importação 2 diz 'three layers of gold'.",
    127: "Estava com o tipo EARRINGS, mas é um colar (Y necklace).",
    81: "Colar personalizável: explique na descrição como a cliente envia o nome (campo de texto ou mensagem).",
    11: "Descrição vazia: escreva comprimento e medidas da pérola.",
    6: "Conjunto montado pela loja: o número no título (003) não é código Rommanel, então fica sem SKU e sem MPN.",
    28: "Conjunto montado pela loja: o número no título (0000002) não é código Rommanel, então fica sem SKU e sem MPN.",
    41: "Conjunto montado pela loja: o número no título (00012026) não é código Rommanel, então fica sem SKU e sem MPN.",
    17: "Descrição vazia: escreva medidas e fecho.",
}

# "Waterproof" promises more than plated jewellery gives; Priscila keeps "water-resistant" (it can get wet)
WATERPROOF = re.compile(r"water\s*-?\s*proof", re.I)

# Product URLs in Portuguese, made only of numbers, or with a typo: change by hand with a redirect (CSV can't change URLs)
HANDLE_FIX = {37, 48, 50, 56, 82, 85, 91, 112, 157, 160}


def slugify(text):
    text = html.unescape(text).lower().replace("é", "e").replace("'", "")
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def strip_html(value):
    text = re.sub(r"<br\s*/?>", " ", value or "")
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", html.unescape(text)).strip()


def product_code(title):
    m = re.search(r"(?:\s-\s*|\s)(\d{3,})\s*$", title)
    return m.group(1) if m else ""


def compose_description(d, title, type_code, rhodium, ring_sizes):
    """The new description: a lead about the piece, its facts as a list and the care line. Returns (html, text)."""
    esc = lambda t: html.escape(t.strip(), quote=False)
    if title.startswith("925"):
        facts = ["Material: solid 925 sterling silver"]
    else:
        facts = ["Material: 18K gold plated, with three layers of gold" + (" and rhodium-plated details" if rhodium else "")]
    size = d.get("size", "").strip()
    if size:
        facts.append(f"{'Length' if type_code in ('N', 'B', 'A') and ' x ' not in size else 'Size'}: {size}")
    if d.get("stones", "").strip():
        facts.append(f"Stones: {d['stones'].strip()}")
    facts += [x.strip() for x in d.get("details", []) if x.strip()]
    if ring_sizes:
        facts.append("Australian ring sizes, with the Brazilian aro in our size guide")
    if type_code == "P" and not any("chain" in f.lower() for f in facts):
        facts.append("Chain sold separately")
    facts += [
        "Hypoallergenic and nickel-free, designed with sensitive skin in mind",
        "Water-resistant and made to last, for everyday wear",
        "Authentic Rommanel, designed in Brazil",
    ]
    care = "Care: put it on after perfume and creams, and take it off for the beach, the pool and hot showers to keep the shine for longer."
    body = f"<p>{esc(d['lead'])}</p>\n<ul>\n" + "".join(f"<li>{esc(f)}</li>\n" for f in facts) + f"</ul>\n<p>{esc(care)}</p>"
    text = d["lead"].strip() + "\n" + "\n".join(f"• {f}" for f in facts) + "\n" + care
    return body, text


def fix_description(body):
    """Removes unsupported claims and typos from a description; returns (new_html, list_of_changes)."""
    changes = []
    new = body
    patterns = [
        (r"\bWater\s*-?\s*proof\b", "Water-resistant", "trocou 'Waterproof' por 'Water-resistant'"),
        (r"\bwater\s*-?\s*proof\b", "water-resistant", "trocou 'waterproof' por 'water-resistant'"),
        (r"\bRomanel\b", "Rommanel", "corrigiu 'Romanel' para 'Rommanel'"),
        (r"Triple gold finish \(24k, 18k &amp; 22k\)|Triple gold finish \(24k, 18k & 22k\)", "Triple-layer gold finish", "trocou '24k, 18k & 22k' por 'Triple-layer gold finish'"),
        (r"\s*Medical grade\.", "", "tirou 'Medical grade'"),
        (r"\bCChain\b", "Chain", "corrigiu 'CChain'"),
    ]
    for pattern, repl, label in patterns:
        new, n = re.subn(pattern, repl, new)
        if n:
            changes.append(label)
    return new, changes


KEEP_CASE = {"Saint", "Benedict", "Our", "Lady", "Aparecida", "Mount", "Carmel", "Holy", "Spirit", "Blessed", "Sacrament",
             "Cinderella", "Lotso", "Fé", "(Faith)", "Y", "18K", "925", "Pet", "Mum", "Best", "Friends"}


def sentence_case(text):
    words = text.split(" ")
    out = [w if w in KEEP_CASE else w.lower() for w in words]
    first = out[0]
    out[0] = first[0].upper() + first[1:]
    return " ".join(out)


def seo_title(title, material_short):
    core = title
    for prefix in ("18K Gold-Plated ", "925 Sterling Silver "):
        core = core.replace(prefix, "")
    core = core.replace(", ", " ")
    candidates = [(f"{core} – {material_short} | Rommanel", 62), (f"{core} – {material_short}", 66),
                  (f"{core} | Rommanel", 62), (core, 70)]
    for text, limit in candidates:
        if len(text) <= limit:
            return text
    return core[:70].rsplit(" ", 1)[0]


def meta_description(title, type_code, tags, sizes):
    silver = title.startswith("925")
    core = title.replace("18K Gold-Plated ", "").replace("925 Sterling Silver ", "").replace(", 2 Pairs", " (2 pairs)")
    length = ""
    m = re.search(r",\s*([0-9.]+ (?:cm|mm))$", core)
    if m:
        core, length = core[: m.start()], m.group(1)
    material = "in solid 925 sterling silver" if silver else "plated in 18K gold"
    if "with Rhodium Plating" in core:
        core = core.replace(" with Rhodium Plating", "")
        material = "plated in 18K gold with rhodium details"
    lead = f"{sentence_case(core)} by Rommanel, {material}"
    if length:
        lead += f", {length} long" if type_code in ("N", "B", "A") else f", {length}"
    parts = [lead + "."]
    tagset = {t.strip() for t in tags.split(",") if t.strip()}
    if "faith" in tagset:
        parts.append("A meaningful gift of faith.")
    elif "kids" in tagset:
        parts.append("A sweet first-jewellery gift.")
    elif "family" in tagset:
        parts.append("A gift for mum to keep her loved ones close.")
    elif type_code == "P":
        parts.append("Add them to your favourite chains." if core.endswith("Pendants") else "Add it to your favourite chain.")
    elif type_code == "R" and sizes:
        parts.append("Available in Australian ring sizes.")
    parts.append("Nickel-free and hypoallergenic.")
    parts.append("Shipped Australia-wide from Perth.")
    text = " ".join(parts)
    while len(text) > 155 and len(parts) > 2:
        parts.pop(-2 if len(parts) > 3 else -1)
        text = " ".join(parts)
    return text


def live_seo(products):
    """What each product page shows Google today: a custom SEO title/description, or Shopify's defaults."""
    live = json.loads(LIVE.read_text())
    by_handle = {p["handle"]: p for p in products}
    norm = lambda v: re.sub(r"\s+", " ", html.unescape(v or "")).strip()
    out = {}
    for r in live["results"]:
        if r["kind"] != "product":
            continue
        p = by_handle[r["handle"]]
        if r.get("status") != 200:
            out[r["handle"]] = {"known": False}
            continue
        title = norm(r.get("title"))
        if title.endswith(" – PurGrace"):
            title = title[: -len(" – PurGrace")]
        desc = norm(r.get("description"))
        body = strip_html(p.get("body_html"))
        default_desc = not desc or body.startswith(desc.rstrip(".… ").rstrip()) or desc[:60] == body[:60]
        out[r["handle"]] = {
            "known": True,
            "title": title if title != norm(p["title"]) else "",
            "description": "" if default_desc else desc,
        }
    return out


def main():
    data = json.loads(FIXTURE.read_text())
    products = data["products"]
    seo_now = live_seo(products)
    membership = {h: set(m) for h, m in data["membership"].items()}
    assert len(products) == len(CURATED), (len(products), len(CURATED))
    written = {d["n"]: d for d in json.loads(DESCRIPTIONS.read_text())} if DESCRIPTIONS.exists() else {}

    rows = []
    for i, (p, (key, title, type_code, theme)) in enumerate(zip(products, CURATED)):
        assert key in p["handle"], (i, key, p["handle"])
        type_name, category, collection = TYPES[type_code]
        silver = title.startswith("925")
        body = p.get("body_html") or ""
        specifics = strip_html(body)
        code = product_code(p["title"])
        rhodium = "rhodium" in (p["title"] + specifics).lower()
        material = "Sterling silver" if silver else "Gold-plated" + ("; Rhodium-plated" if rhodium else "")
        sizes = [v["option1"] for v in p["variants"]] if len(p["variants"]) > 1 else []
        stone_tags = []
        lower = (title + " " + specifics).lower()
        for word, tag in (("zirconi", "cubic zirconia"), ("crystal", "crystal"), ("pearl", "pearl"), ("resin", "resin")):
            if word in lower and tag not in stone_tags:
                stone_tags.append(tag)
        tags = ["925 sterling silver" if silver else "18k gold plated"] + [t.strip() for t in theme.split(",") if t.strip()] + stone_tags
        tags = list(dict.fromkeys(tags))
        material_short = "Sterling Silver" if silver else "18K Gold Plated"
        should_be_in = {c for c in [collection] if c}
        should_be_in.add("solid-925-silver" if silver else "18k-gold-plated")
        if "faith" in tags:
            should_be_in.add("faith")
        if "kids" in tags:
            should_be_in.add("kids")
        missing = sorted(c for c in should_be_in if p["handle"] not in membership.get(c, set()))
        wrong = sorted(c for c in ("18k-gold-plated", "solid-925-silver") if p["handle"] in membership.get(c, set()) and c not in should_be_in)
        new_body, changes = fix_description(body)
        new_text = ""
        if i in written:
            new_body, new_text = compose_description(written[i], title, type_code, rhodium, sizes)
            changes = ["descrição nova"] + [c for c in changes if "Romanel" not in c]
        rows.append(
            {
                "n": i,
                "handle": p["handle"],
                "url": f"{STORE}/products/{p['handle']}",
                "old_title": p["title"],
                "title": title,
                "code": code,
                "old_type": p.get("product_type") or "",
                "type": type_name,
                "category": category,
                "material": material,
                "tags": ", ".join(tags),
                "seo_title": seo_title(title, material_short),
                "seo_description": meta_description(title, type_code, ", ".join(tags), sizes),
                "seo_known": seo_now.get(p["handle"], {}).get("known", False),
                "seo_title_now": seo_now.get(p["handle"], {}).get("title", ""),
                "seo_description_now": seo_now.get(p["handle"], {}).get("description", ""),
                "is_set": type_code == "S",
                "category_collection": collection or "",
                "alt": title,
                "missing": ", ".join(missing),
                "wrong": ", ".join(wrong),
                "sizes": ", ".join(sizes),
                "variants": p["variants"],
                "options": p["options"],
                "old_skus": [v.get("sku") or "" for v in p["variants"]],
                "water": bool(WATERPROOF.search(body)),
                "body": body,
                "new_body": new_body,
                "new_text": new_text,
                "changes": changes,
                "flag": FLAGS.get(i, ""),
                "new_handle": f"{slugify(title)}-{code}" if i in HANDLE_FIX else "",
                "price": p["variants"][0]["price"],
                "images": len(p["images"]),
            }
        )

    dupes = Counter(r["title"] for r in rows)
    for r in rows:
        if dupes[r["title"]] > 1 and not r["flag"]:
            r["flag"] = "Título repetido em outro produto."
    # Her own SEO titles: typos and titles shared by two products
    now_dupes = Counter(r["seo_title_now"] for r in rows if r["seo_title_now"])
    for r in rows:
        notes = []
        if re.match(r"^\d?K ", r["seo_title_now"]) and not r["seo_title_now"].startswith("18K"):
            notes.append(f"O título para o Google de hoje começa com \"{r['seo_title_now'][:3]}\": falta o 1 de 18K.")
        if r["seo_title_now"] and now_dupes[r["seo_title_now"]] > 1:
            notes.append("O título para o Google de hoje é igual ao de outro produto: diferencie (tamanho, pedra, cor).")
        if notes:
            r["flag"] = " ".join(filter(None, [r["flag"]] + notes))
        # Import 3 only fills products that have neither text today, so nothing she wrote is replaced
        r["seo_fill"] = r["seo_known"] and not r["seo_title_now"] and not r["seo_description_now"]
    seo_dupes = Counter(r["seo_title"] for r in rows)
    write_csv_import(rows)
    write_descriptions_import(rows)
    write_seo_import(rows)
    write_undo(rows)
    write_workbook(rows, data, dupes, seo_dupes)
    print(f"{len(rows)} products · {sum(1 for r in rows if r['changes'])} descriptions fixed · "
          f"{sum(1 for r in rows if r['seo_title_now'])} custom SEO titles kept · {sum(1 for r in rows if r['seo_fill'])} products in import 3 · "
          f"{sum(1 for r in rows if r['missing'])} missing from a collection · "
          f"duplicate titles {sum(1 for v in dupes.values() if v > 1)}")


def variant_rows(r, first):
    """One CSV row per variant, keeping the option names and values exactly as they are so no variant is recreated."""
    out = []
    names = [o["name"] for o in r["options"]]
    for k, v in enumerate(r["variants"]):
        row = {"URL handle": r["handle"]}
        if k == 0:
            row.update(first)
        for n, name in enumerate(names, start=1):
            if k == 0:
                row[f"Option{n} name"] = name
            row[f"Option{n} value"] = v.get(f"option{n}") or ""
        out.append(row)
    return out


# Two products for a first test import: a single piece and a ring with sizes
TEST_PRODUCTS = (0, 122)
MPN = "Google Shopping / Manufacturer part number (MPN)"


def open_csv(name, headers):
    fh = open(HERE / name, "w", newline="", encoding="utf-8")
    writer = csv.DictWriter(fh, fieldnames=headers, lineterminator="\n")
    writer.writeheader()
    return fh, writer


def sku_for(r, k):
    """The Rommanel reference: SKU for stock and MPN for Google. Ring sizes keep what they have;
    the sets are put together by the shop, so they get no manufacturer code."""
    existing = r["old_skus"][k]
    if existing or len(r["variants"]) > 1 or r["is_set"]:
        return existing
    return r["code"]


def write_csv_import(rows):
    write_import_file(rows, "shopify-import-1-titulos-tipos-tags.csv")
    write_import_file([rows[i] for i in TEST_PRODUCTS], "shopify-import-0-teste-2-produtos.csv")


def write_import_file(rows, name):
    headers = ["URL handle", "Title", "Type", "Product category", "Tags", "Option1 name", "Option1 value", "SKU", MPN]
    fh, w = open_csv(name, headers)
    with fh:
        for r in rows:
            first = {"Title": r["title"], "Type": r["type"], "Product category": r["category"], "Tags": r["tags"]}
            for k, row in enumerate(variant_rows(r, first)):
                row["SKU"] = sku_for(r, k)
                row[MPN] = "" if r["is_set"] or len(r["variants"]) > 1 else sku_for(r, k)
                w.writerow(row)


def write_descriptions_import(rows):
    fh, w = open_csv("shopify-import-2-descricoes.csv", ["URL handle", "Title", "Description", "Option1 name", "Option1 value"])
    with fh:
        for r in rows:
            if not r["changes"]:
                continue
            assert not WATERPROOF.search(r["new_body"]), r["handle"]
            for row in variant_rows(r, {"Title": r["title"], "Description": r["new_body"]}):
                w.writerow(row)


def write_seo_import(rows):
    fh, w = open_csv("shopify-import-3-seo-onde-falta.csv", ["URL handle", "Title", "SEO title", "SEO description", "Option1 name", "Option1 value"])
    with fh:
        for r in rows:
            if r["seo_fill"]:
                for row in variant_rows(r, {"Title": r["title"], "SEO title": r["seo_title"], "SEO description": r["seo_description"]}):
                    w.writerow(row)


def write_undo(rows):
    """Puts back what imports 1 and 2 change, as it was on 5/10/2026 (the Shopify category can't be put back)."""
    fh, w = open_csv("shopify-import-desfazer.csv", ["URL handle", "Title", "Type", "Tags", "Description", "Option1 name", "Option1 value", "SKU", MPN])
    with fh:
        for r in rows:
            first = {"Title": r["old_title"], "Type": r["old_type"], "Tags": "", "Description": r["body"]}
            for k, row in enumerate(variant_rows(r, first)):
                row["SKU"] = r["old_skus"][k]
                row[MPN] = ""
                w.writerow(row)


HEAD_FILL = PatternFill("solid", fgColor="2E211E")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FBE3DC")
NEW_FILL = PatternFill("solid", fgColor="EEF6EE")


def sheet(wb, name, headers, widths, rows, warn_col=None, new_cols=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if c in new_cols:
                cell.fill = NEW_FILL
        if warn_col and ws.cell(row=r, column=warn_col).value:
            ws.cell(row=r, column=warn_col).fill = WARN_FILL
    ws.freeze_panes = "C2" if name == "Produtos" else "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    return ws


def write_workbook(rows, data, dupes, seo_dupes):
    wb = Workbook()
    readme = wb.active
    readme.title = "Leia-me"
    lines = [
        ("PurGrace · revisão de SEO dos produtos", True),
        ("", False),
        ("O que tem aqui", True),
        ("Produtos: título limpo (sem código e sem CAIXA ALTA), tipo, categoria do Shopify, material, tags, código Rommanel, o título e a descrição para o Google que a loja mostra hoje ao lado de uma sugestão, texto alternativo da foto principal e o que falta em cada produto. Colunas verdes = sugestão nova.", False),
        ("Coleções: o texto que cada coleção tem hoje, o que tem de errado nele, e o título, o texto de abertura e os textos para o Google sugeridos.", False),
        ("Problemas: tudo que precisa de uma decisão sua (anúncios repetidos, códigos estranhos, tamanhos de anel, frases sem comprovação).", False),
        ("Anéis: tamanhos oferecidos hoje e o aro que o código Rommanel indica.", False),
        ("", False),
        ("Como aplicar (o passo a passo completo está no guia de SEO)", True),
        ("Faça as importações antes de mexer à mão em qualquer produto (links, tamanhos, títulos, descrições ou apagar anúncio). Os arquivos acham cada produto pelo link e pelos tamanhos de 5/10/2026.", False),
        ("1. Backup: Produtos › Exportar › Todos os produtos › Arquivo CSV simples (Plain CSV file). O arquivo chega por e-mail: baixe e guarde. Não importe esse backup inteiro depois: ele volta o estoque e as fotos para o dia do backup.", False),
        ("2. Em Produtos › Importar, escolha o arquivo, desmarque 'Publicar novos produtos em todos os canais de vendas' e marque 'Sobrescrever produtos com identificadores correspondentes' (Overwrite products with matching handles). No resumo, confira o número de produtos antes de clicar em Importar.", False),
        ("3. Importe nesta ordem, esperando o e-mail de confirmação de cada um: shopify-import-0-teste-2-produtos.csv (2 produtos; confira na loja), shopify-import-1-titulos-tipos-tags.csv (162), shopify-import-2-descricoes.csv (162 descrições novas; leia algumas na coluna 'Descrição nova' antes) e, se quiser, shopify-import-3-seo-onde-falta.csv (só os produtos sem título e descrição para o Google hoje).", False),
        ("Se algo sair errado, importe shopify-import-desfazer.csv do mesmo jeito: ele volta títulos, tipos, tags, SKU e descrições para o que eram em 5/10/2026.", False),
        ("Os títulos e descrições para o Google que você já escreveu não mudam: a importação 1 não tem essas colunas, e a 3 só preenche os produtos que estão sem.", False),
        ("Os links (URLs), preços, fotos e estoque não mudam. As colunas Option1 name/Option1 value estão iguais às de hoje para o Shopify não recriar os tamanhos dos anéis: não apague essas colunas.", False),
        ("Não abra os CSV no Excel para salvar de novo: ele estraga acentos e travessões.", False),
    ]
    for text, bold in lines:
        readme.append([text])
        readme.cell(row=readme.max_row, column=1).font = Font(bold=bold, size=13 if bold else 11)
        readme.cell(row=readme.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 120

    headers = ["#", "Link", "Título atual", "Título novo", "Código Rommanel (SKU e MPN)", "Tipo atual", "Tipo novo",
               "Categoria do Shopify", "Material (metacampo da categoria)", "Tags",
               "Título para o Google hoje", "Título sugerido", "Car.",
               "Descrição para o Google hoje", "Descrição sugerida", "Car.", "Importação 3 preenche?",
               "Texto alternativo da foto principal", "Falta nas coleções", "Está na coleção errada", "Tamanhos (variantes)",
               "Descrição será corrigida", "Descrição nova (importação 2)", "URL nova sugerida (mudar à mão, com redirecionamento)", "Atenção"]
    widths = [5, 10, 40, 44, 14, 11, 12, 28, 18, 30, 40, 40, 6, 50, 50, 6, 12, 40, 22, 18, 18, 30, 70, 34, 50]
    body = []
    for r in rows:
        title_now = r["seo_title_now"] or ("(sem: usa o título do produto)" if r["seo_known"] else "(não deu para ler)")
        desc_now = r["seo_description_now"] or ("(sem: usa o começo da descrição)" if r["seo_known"] else "(não deu para ler)")
        body.append([
            r["n"], r["url"], r["old_title"], r["title"], "" if r["is_set"] else r["code"], r["old_type"], r["type"], r["category"],
            r["material"], r["tags"], title_now, r["seo_title"], len(r["seo_title"]), desc_now, r["seo_description"],
            len(r["seo_description"]), "sim" if r["seo_fill"] else "não", r["alt"], r["missing"], r["wrong"], r["sizes"],
            "; ".join(r["changes"]), r["new_text"], r["new_handle"], r["flag"],
        ])
    ws = sheet(wb, "Produtos", headers, widths, body, warn_col=25, new_cols=(4, 5, 7, 8, 9, 10, 12, 15, 18, 23, 24))
    for row in range(2, ws.max_row + 1):
        link = ws.cell(row=row, column=2)
        link.hyperlink, link.value, link.font = link.value, "abrir", Font(color="A65A4F", underline="single")

    # Collections
    membership = data["membership"]
    by_type = defaultdict(int)
    missing_count = Counter()
    for r in rows:
        by_type[r["type"]] += 1
        for c in filter(None, (x.strip() for x in r["missing"].split(","))):
            missing_count[c] += 1
    wrong_gold = sum(1 for r in rows if "18k-gold-plated" in r["wrong"])
    collections = [
        ("earrings", "Earrings", "Gold Plated Earrings – Hoops, Studs & Ear Cuffs | Rommanel",
         "Rommanel earrings plated in 18K gold or in 925 sterling silver: hoops, studs, ear cuffs and kids' earrings. Nickel-free, shipped from Perth Australia-wide.",
         "Hoops for every day, studs that sparkle and ear cuffs that need no piercing. Every pair is authentic Rommanel, plated in 18K gold or made in 925 sterling silver, nickel-free and hypoallergenic."),
        ("necklace", "Necklaces", "Gold Plated Necklaces & Chains | Rommanel Australia",
         "18K gold-plated and 925 silver necklaces by Rommanel: fine chains, pendant necklaces and scapulars from 35 to 80 cm. Shipped from Perth Australia-wide.",
         "Fine chains to wear alone or layer, pendant necklaces and scapulars. Each listing shows the length in centimetres, so you can pick the right drop for your neckline."),
        ("pendant", "Pendants", "Gold Plated Pendants, Medals & Charms | Rommanel Australia",
         "Rommanel pendants and medals plated in 18K gold: saints, hearts, butterflies and family charms. Add one to your favourite chain. Shipped from Perth.",
         "Saints and medals, hearts, butterflies and little boys and girls to carry your family with you. Pendants come without a chain, so you can pair them with one you love."),
        ("ring", "Rings", "Gold Plated & Silver Rings | Rommanel Australia",
         "Rommanel rings plated in 18K gold and in 925 sterling silver, in Australian sizes. Use our size guide to find your fit. Shipped from Perth.",
         "Solitaires, half eternity bands and statement rings. Our size guide converts Australian letter sizes to the Brazilian aro used by Rommanel."),
        ("bracelet", "Bracelets", "Gold Plated Bracelets & Anklets | Rommanel Australia",
         "18K gold-plated and 925 silver bracelets by Rommanel: hearts, butterflies, beads and chains, most adjustable. Shipped from Perth Australia-wide.",
         "Delicate chains, hearts and butterflies, most with an adjustable length. Wear one alone or stack a few."),
        ("set", "Jewellery Sets", "Gold Plated Jewellery Sets – Necklace & Earrings | Rommanel",
         "Matching Rommanel necklace and earring sets plated in 18K gold. A ready-made gift, shipped from Perth Australia-wide.",
         "Necklace and earrings that match, ready to give."),
        ("faith", "Faith Jewellery", "Catholic Jewellery – Saint Medals, Crosses & Scapulars | Rommanel",
         "Saint Benedict medals, Our Lady of Aparecida, Holy Spirit pendants, crosses and scapulars plated in 18K gold by Rommanel. Shipped from Perth.",
         "Carry your faith close to your heart: Saint Benedict, Our Lady of Aparecida and of Mount Carmel, the Holy Spirit, crosses and scapulars."),
        ("kids", "Kids' Jewellery", "Kids' Gold Plated Earrings & Bracelets | Rommanel Australia",
         "Small Rommanel earrings and bracelets for children, plated in 18K gold and nickel-free. A sweet first-jewellery gift, shipped from Perth.",
         "Little studs with screw backs, small bracelets and playful charms for children."),
        ("solid-925-silver", "925 Sterling Silver", "925 Sterling Silver Jewellery | Rommanel Australia",
         "Solid 925 sterling silver earrings, pendants, rings and chains by Rommanel, many with cubic zirconia. Shipped from Perth Australia-wide.",
         "Solid 925 sterling silver for when you want a cooler shine."),
        ("18k-gold-plated", "18K Gold Plated", "18K Gold Plated Jewellery | Rommanel Australia",
         "Brazilian Rommanel jewellery plated in 18K gold: earrings, necklaces, pendants, rings and bracelets. Nickel-free, shipped from Perth.",
         "Rommanel has been plating jewellery in Brazil since 1986. Every piece here is authentic Rommanel, plated in 18K gold."),
        ("all", "All Jewellery", "Rommanel Jewellery Australia – 18K Gold Plated & 925 Silver",
         "Shop authentic Rommanel jewellery in Australia: 18K gold-plated and 925 sterling silver earrings, necklaces, rings and more, shipped from Perth.",
         "Every piece here is authentic Rommanel jewellery from Brazil, plated in 18K gold or made in solid 925 sterling silver, nickel-free and chosen in Perth."),
        ("best-sellers", "Best Sellers", "Best-Selling Rommanel Jewellery | PurGrace",
         "The Rommanel pieces our customers come back for: hoops, butterflies, saints and fine chains, plated in 18K gold. Shipped from Perth.",
         ""),
        ("mother-s-day", "Mother's Day Gifts", "Mother's Day Jewellery Gifts | Rommanel Australia",
         "Mother's Day gifts by Rommanel: boy and girl pendants, hearts and family bracelets plated in 18K gold. Shipped from Perth Australia-wide.",
         "Gifts she'll wear every day: boy and girl pendants, hearts and family bracelets, plated in 18K gold by Rommanel."),
        ("eofy-gold", "EOFY Gold", "EOFY Sale – 18K Gold Plated Jewellery | Rommanel Australia",
         "Selected 18K gold-plated Rommanel jewellery on sale for the end of the financial year. Nickel-free, shipped from Perth Australia-wide.",
         "Selected 18K gold-plated Rommanel pieces for the end of the financial year, nickel-free and hypoallergenic."),
        ("gifts-for-her", "Gifts for Her", "Gifts for Her – Hearts, Butterflies & Sets | Rommanel",
         "Rommanel gifts for her: heart and butterfly jewellery and ready-to-give sets plated in 18K gold, for a girlfriend, wife or friend. Shipped from Perth.",
         "Hearts, butterflies and ready-to-give sets for your girlfriend, your wife or your best friend."),
        ("gifts-for-him", "Gifts for Him", "Gifts for Him – Chains, Crosses & Scapulars | Rommanel",
         "18K gold-plated chains, crosses and scapulars by Rommanel for your boyfriend, husband or dad. Nickel-free, shipped from Perth Australia-wide.",
         "Chains, crosses and scapulars for your boyfriend, your husband or your dad."),
        ("disney", "Disney", "Disney Jewellery – 18K Gold Plated | Rommanel Australia",
         "Pieces from the Rommanel Disney line plated in 18K gold, like Cinderella and Lotso: a magical gift for kids and fans. Shipped from Perth.",
         "Pieces from the Rommanel Disney line, plated in 18K gold."),
    ]
    # Gift collections that don't exist yet: smart collections by the tags import 1 adds
    planned = {
        "gifts-for-her": ("love", "butterfly", "gift set"),
        "gifts-for-him": ("men", "unisex"),
        "disney": ("disney",),
    }
    notes = {
        "mother-s-day": "Deixe a página no ar o ano todo e troque os produtos em abril, sem mudar o endereço: assim ela não perde a posição no Google.",
        "eofy-gold": "A promoção de fim de ano fiscal foi em junho. Se não houver desconto agora, tire o '30% OFF' do texto ou esconda a coleção do menu até a próxima.",
        "best-sellers": "O texto de hoje pode ficar.",
        "gifts-for-her": "Coleção nova para o guia de presentes. Crie como automática, com qualquer condição: Tag é igual a love, Tag é igual a butterfly, Tag é igual a gift set.",
        "gifts-for-him": "Coleção nova para o guia de presentes. Crie como automática, com qualquer condição: Tag é igual a men, Tag é igual a unisex.",
        "disney": "Coleção nova para o guia de presentes. Crie como automática: Tag é igual a disney. Peça nova da linha Disney entra sozinha quando tiver a tag disney.",
    }
    claims = [
        (WATERPROOF, "diz 'waterproof'; troque por 'water-resistant'"),
        (re.compile(r"24k|22k", re.I), "diz '24k, 18k e 22k', que a Rommanel não publica"),
        (re.compile(r"30% off", re.I), "anuncia 30% OFF"),
    ]
    live = {r["handle"]: r for r in json.loads(LIVE.read_text())["results"] if r["kind"] == "collection"}
    generic = Counter(re.sub(r"\s+", " ", (r.get("visibleDescription") or "")).strip() for r in live.values())
    by_handle = {c["handle"]: c for c in data["collections"]}
    coll_rows = []
    for handle, title, stitle, sdesc, intro in collections:
        c = by_handle.get(handle, {})
        now = re.sub(r"\s+", " ", (live.get(handle, {}).get("visibleDescription") or "")).strip()
        problems = [label for rx, label in claims if rx.search(now)]
        if now and generic[now] > 1:
            problems.append(f"texto igual ao de outras {generic[now] - 1} coleções")
        attention = [notes.get(handle, "")]
        if handle == "18k-gold-plated" and wrong_gold:
            attention.append(f"{wrong_gold} peças de prata estão nesta coleção.")
        coll_rows.append([
            handle, f"{STORE}/collections/{handle}", c.get("title", ""), title, now, "; ".join(problems),
            intro or "(mantenha o texto de hoje)", stitle, len(stitle), sdesc, len(sdesc),
            len(membership.get(handle, [])),
            sum(1 for r in rows if set(planned[handle]) & {t.strip() for t in r["tags"].split(",")}) if handle in planned else missing_count.get(handle, 0),
            " ".join(filter(None, attention)),
        ])
    ws = sheet(wb, "Coleções", ["Coleção", "Link", "Título atual", "Título novo", "Texto de hoje (na loja)", "O que tem de errado",
                                "Texto de abertura novo", "Título para o Google", "Car.", "Descrição para o Google", "Car.",
                                "Produtos hoje", "Produtos faltando", "Atenção"],
               [18, 10, 18, 20, 50, 30, 55, 46, 6, 55, 6, 10, 10, 40], coll_rows, warn_col=6, new_cols=(4, 7, 8, 10))
    for row in range(2, ws.max_row + 1):
        link = ws.cell(row=row, column=2)
        link.hyperlink, link.value, link.font = link.value, "abrir", Font(color="A65A4F", underline="single")

    # Issues
    issues = []
    for r in rows:
        if r["flag"]:
            issues.append([r["n"], r["title"], r["url"], r["flag"]])
    for r in rows:
        if r["new_handle"]:
            issues.append([r["n"], r["title"], r["url"], f"O link do produto está em português, só com números ou com erro ({r['handle']}). No produto, em 'Search engine listing' (listagem nos mecanismos de pesquisa), troque o 'URL handle' para {r['new_handle']} e deixe marcado 'Create a URL redirect' (criar redirecionamento)."])
    for r in rows:
        if r["water"]:
            fix = "A descrição nova da importação 2 diz 'water-resistant'" if r["new_text"] else "A importação 2 troca por 'Water-resistant'"
            issues.append([r["n"], r["title"], r["url"], f"A descrição de hoje diz 'Waterproof' (à prova d'água). {fix}, que diz que pode molhar sem prometer que aguenta mar e piscina."])
    issues.append(["–", "Todas as fotos", "", "Nenhuma das 436 fotos tem texto alternativo. O tema novo usa o título limpo como texto alternativo enquanto você não escreve um, mas o ideal é descrever cada foto (ex.: 'brinco de borboleta na orelha da modelo')."])
    spellings = sorted({r["old_type"] for r in rows if r["old_type"]})
    blank_types = sum(1 for r in rows if not r["old_type"])
    issues.append(["–", "Tipos de produto", "", f"Hoje há {len(spellings)} grafias diferentes para as mesmas coisas (EARRINGS, EARRING, Brinco, CHEKER...) e {blank_types} produtos sem tipo. A importação 1 deixa só 7: Earrings, Necklaces, Pendants, Rings, Bracelets, Anklets, Jewellery Sets."])
    outside = [r for r in rows if r["category_collection"] and r["category_collection"] in r["missing"].split(", ")]
    earrings_out = sum(1 for r in outside if r["category_collection"] == "earrings")
    issues.append(["–", "Coleções incompletas", "", f"{len(outside)} produtos não estão na coleção da categoria deles (ex.: {earrings_out} brincos fora de 'Earrings'). Com os tipos padronizados, dá para transformar as coleções em automáticas (veja o guia)."])
    live_all = json.loads(LIVE.read_text())["results"]
    for r in live_all:
        text = (r.get("visibleDescription") or "") + " " + (r.get("description") or "")
        found = []
        if WATERPROOF.search(text):
            found.append("diz 'waterproof' (troque por 'water-resistant')")
        if re.search(r"24k|22k", text, re.I):
            found.append("diz '24k, 18k e 22k' (a Rommanel só diz 'três camadas de ouro')")
        if not found or r["kind"] not in ("collection", "page"):
            continue
        where = "Produtos › Coleções" if r["kind"] == "collection" else "Loja virtual › Páginas"
        name = "Coleção " if r["kind"] == "collection" else "Página "
        issues.append(["–", name + r.get("title", "").replace(" – PurGrace", ""), r["url"], f"O texto {' e '.join(found)}. Corrija em {where}, no campo Descrição/Conteúdo; a importação não mexe nisso."])
    ws = sheet(wb, "Problemas", ["#", "Produto", "Link", "O que fazer"], [5, 44, 12, 100], issues)
    for row in range(2, ws.max_row + 1):
        link = ws.cell(row=row, column=3)
        if link.value:
            link.hyperlink, link.value, link.font = link.value, "abrir", Font(color="A65A4F", underline="single")

    # Rings
    ring_rows = []
    for r in rows:
        if r["type"] != "Rings":
            continue
        code = r["code"]
        aro = code[6:8] if len(code) == 10 else ""
        ring_rows.append([r["n"], r["title"], code, aro or "–", r["sizes"] or "(sem opção de tamanho)", r["flag"]])
    sheet(wb, "Anéis", ["#", "Anel", "Código", "Aro pelo código (dígitos 7-8)", "Tamanhos oferecidos hoje", "Atenção"],
          [5, 50, 14, 16, 34, 70], ring_rows, warn_col=6)

    wb.save(HERE / "PurGrace-SEO.xlsx")


if __name__ == "__main__":
    main()
